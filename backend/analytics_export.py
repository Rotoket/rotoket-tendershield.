"""
Экспорт аналитики в CSV и Excel
"""

import logging
import csv
from io import BytesIO, StringIO
from typing import Dict, List, Optional
from datetime import datetime
from analytics import AnalyticsService

logger = logging.getLogger(__name__)


def export_analytics_to_csv(
    analytics_service: AnalyticsService,
    user_id: Optional[int] = None,
    days: int = 30
) -> BytesIO:
    """
    Экспортирует аналитику в CSV
    
    Args:
        analytics_service: Экземпляр AnalyticsService
        user_id: ID пользователя (если None - системная статистика)
        days: Количество дней для временной линии
        
    Returns:
        BytesIO буфер с CSV данными
    """
    output = StringIO()
    writer = csv.writer(output)
    
    # Заголовок
    writer.writerow(['Отчет по аналитике', f'Сгенерирован: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'])
    writer.writerow([])
    
    if user_id:
        # Статистика пользователя
        stats = analytics_service.get_user_stats(user_id)
        writer.writerow(['Статистика пользователя'])
        writer.writerow(['Email', stats['email']])
        writer.writerow(['Тариф', stats['plan_type']])
        writer.writerow(['Всего анализов', stats['total_analyses']])
        writer.writerow(['Пакетных анализов', stats['total_packages']])
        writer.writerow(['Анализов за месяц', stats['current_month_analyses']])
        writer.writerow(['Лимит тарифа', stats['tariff_limit']])
        writer.writerow(['Триал активен', 'Да' if stats['trial_active'] else 'Нет'])
        if stats['trial_active']:
            writer.writerow(['Дней триала осталось', stats['trial_days_left']])
        writer.writerow([])
    
    # Временная линия
    timeline = analytics_service.get_analyses_timeline(user_id, days)
    writer.writerow(['Временная линия анализов'])
    writer.writerow(['Дата', 'Количество'])
    for item in timeline:
        writer.writerow([item['date'], item['count']])
    writer.writerow([])
    
    # Популярные отрасли
    industries = analytics_service.get_popular_industries(10)
    writer.writerow(['Популярные отрасли'])
    writer.writerow(['Отрасль', 'Количество'])
    for industry in industries:
        writer.writerow([industry['industry'], industry['count']])
    writer.writerow([])
    
    # Средние оценки
    scores = analytics_service.get_average_scores()
    writer.writerow(['Средние оценки по отраслям'])
    writer.writerow(['Отрасль', 'Средняя оценка', 'Количество анализов'])
    for industry, data in scores.items():
        writer.writerow([industry, f"{data['average_score']:.2f}", data['analyses_count']])
    
    # Конвертируем в BytesIO
    output.seek(0)
    csv_bytes = BytesIO()
    csv_bytes.write(output.getvalue().encode('utf-8-sig'))  # UTF-8 с BOM для Excel
    csv_bytes.seek(0)
    
    return csv_bytes


def export_analytics_to_excel(
    analytics_service: AnalyticsService,
    user_id: Optional[int] = None,
    days: int = 30
) -> BytesIO:
    """
    Экспортирует аналитику в Excel
    
    Args:
        analytics_service: Экземпляр AnalyticsService
        user_id: ID пользователя (если None - системная статистика)
        days: Количество дней для временной линии
        
    Returns:
        BytesIO буфер с Excel файлом
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise ImportError("openpyxl не установлен. Установите: pip install openpyxl")
    
    wb = openpyxl.Workbook()
    
    # Стили
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    title_font = Font(bold=True, size=14)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Лист 1: Общая статистика
    ws = wb.active
    ws.title = "Статистика"
    row = 1
    
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'A{row}'] = f"Отчет по аналитике - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws[f'A{row}'].font = title_font
    row += 2
    
    if user_id:
        stats = analytics_service.get_user_stats(user_id)
        ws[f'A{row}'] = "Статистика пользователя"
        ws[f'A{row}'].font = title_font
        row += 1
        
        user_data = [
            ["Email", stats['email']],
            ["Тариф", stats['plan_type']],
            ["Всего анализов", stats['total_analyses']],
            ["Пакетных анализов", stats['total_packages']],
            ["Анализов за месяц", stats['current_month_analyses']],
            ["Лимит тарифа", stats['tariff_limit']],
        ]
        
        for data_row in user_data:
            ws[f'A{row}'] = data_row[0]
            ws[f'B{row}'] = data_row[1]
            ws[f'A{row}'].font = Font(bold=True)
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        row += 1
    
    # Лист 2: Временная линия
    ws2 = wb.create_sheet("Временная линия")
    row = 1
    ws2[f'A{row}'] = "Дата"
    ws2[f'B{row}'] = "Количество"
    ws2[f'A{row}'].fill = header_fill
    ws2[f'A{row}'].font = header_font
    ws2[f'B{row}'].fill = header_fill
    ws2[f'B{row}'].font = header_font
    row += 1
    
    timeline = analytics_service.get_analyses_timeline(user_id, days)
    for item in timeline:
        ws2[f'A{row}'] = item['date']
        ws2[f'B{row}'] = item['count']
        ws2[f'A{row}'].border = border
        ws2[f'B{row}'].border = border
        row += 1
    
    # Лист 3: Отрасли
    ws3 = wb.create_sheet("Отрасли")
    row = 1
    ws3[f'A{row}'] = "Отрасль"
    ws3[f'B{row}'] = "Количество"
    ws3[f'A{row}'].fill = header_fill
    ws3[f'A{row}'].font = header_font
    ws3[f'B{row}'].fill = header_fill
    ws3[f'B{row}'].font = header_font
    row += 1
    
    industries = analytics_service.get_popular_industries(20)
    for industry in industries:
        ws3[f'A{row}'] = industry['industry']
        ws3[f'B{row}'] = industry['count']
        ws3[f'A{row}'].border = border
        ws3[f'B{row}'].border = border
        row += 1
    
    # Настройка ширины колонок
    for ws in [wb.active, ws2, ws3]:
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
    
    # Сохранение в буфер
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return buffer

