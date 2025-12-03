"""
Экспорт результатов анализа в Excel
"""

import logging
from io import BytesIO
from typing import Dict, Any, List
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def generate_excel_report(analysis_data: Dict[str, Any]) -> BytesIO:
    """
    Генерирует Excel отчет на основе данных анализа
    
    Args:
        analysis_data: Данные анализа
        
    Returns:
        BytesIO буфер с Excel файлом
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        raise ImportError("openpyxl не установлен. Установите: pip install openpyxl")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Анализ тендера"
    
    # Стили
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    title_font = Font(bold=True, size=14)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    row = 1
    
    # Заголовок
    ws.merge_cells(f'A{row}:D{row}')
    ws[f'A{row}'] = "Отчет об анализе тендерной документации"
    ws[f'A{row}'].font = title_font
    ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
    row += 2
    
    # Общая информация
    info_data = [
        ["Параметр", "Значение"],
        ["Файл", analysis_data.get("filename", "N/A")],
        ["Отрасль", analysis_data.get("industry", "UNIVERSAL")],
        ["Дата анализа", analysis_data.get("created_at", datetime.utcnow()).strftime("%Y-%m-%d %H:%M") if isinstance(analysis_data.get("created_at"), datetime) else str(analysis_data.get("created_at", "N/A"))],
        ["Общая оценка", f"{analysis_data.get('score', 'N/A')} / 100"],
        ["Вердикт", analysis_data.get("verdict", "N/A")],
    ]
    
    for info_row in info_data:
        ws[f'A{row}'] = info_row[0]
        ws[f'B{row}'] = info_row[1]
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        row += 1
    
    row += 1
    
    # Краткое резюме
    ws[f'A{row}'] = "Краткое резюме:"
    ws[f'A{row}'].font = Font(bold=True)
    row += 1
    ws.merge_cells(f'A{row}:D{row}')
    ws[f'A{row}'] = analysis_data.get("summary", "Нет описания.")
    ws[f'A{row}'].alignment = Alignment(wrap_text=True, vertical='top')
    row += 2
    
    # Паспорт тендера
    ws[f'A{row}'] = "Паспорт тендера"
    ws[f'A{row}'].font = title_font
    row += 1
    
    passport = analysis_data.get("passport", {})
    passport_data = [
        ["Параметр", "Значение"],
        ["НМЦК", passport.get("nmck", "Не указано")],
        ["Регион", passport.get("region", "Не указано")],
        ["Закон (ФЗ)", passport.get("fz", "Не указано")],
        ["Срок подачи заявки", passport.get("deadlineApp", "Не указано")],
        ["Обеспечение", passport.get("guarantee", "Не указано")],
    ]
    
    for i, passport_row in enumerate(passport_data):
        ws[f'A{row}'] = passport_row[0]
        ws[f'B{row}'] = passport_row[1]
        if i == 0:  # Заголовок
            ws[f'A{row}'].fill = header_fill
            ws[f'A{row}'].font = header_font
            ws[f'B{row}'].fill = header_fill
            ws[f'B{row}'].font = header_font
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        row += 1
    
    row += 1
    
    # Выявленные риски
    issues = analysis_data.get("issues", [])
    if issues:
        ws[f'A{row}'] = "Выявленные риски"
        ws[f'A{row}'].font = title_font
        row += 1
        
        # Заголовки таблицы рисков
        headers = ["Название", "Серьезность", "Описание", "Цитата"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        row += 1
        
        # Данные рисков
        for issue in issues:
            severity = issue.get("severity", "LOW")
            severity_color = {
                "HIGH": "FF4444",
                "MEDIUM": "FFA500",
                "LOW": "90EE90"
            }.get(severity, "FFFFFF")
            
            ws.cell(row=row, column=1, value=issue.get("title", "Без названия")).border = border
            severity_cell = ws.cell(row=row, column=2, value=severity)
            severity_cell.fill = PatternFill(start_color=severity_color, end_color=severity_color, fill_type="solid")
            severity_cell.border = border
            ws.cell(row=row, column=3, value=issue.get("description", "Нет описания.")).border = border
            ws.cell(row=row, column=4, value=issue.get("quote", "")).border = border
            
            # Перенос текста для описания и цитаты
            ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical='top')
            ws.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical='top')
            row += 1
    
    row += 1
    
    # Спецификация
    specs = analysis_data.get("specs", [])
    if specs:
        ws[f'A{row}'] = "Спецификация (ключевые позиции)"
        ws[f'A{row}'].font = title_font
        row += 1
        
        # Заголовки
        spec_headers = ["Наименование", "Количество", "Характеристики"]
        for col, header in enumerate(spec_headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        row += 1
        
        # Данные спецификации
        for spec in specs:
            ws.cell(row=row, column=1, value=spec.get("name", "N/A")).border = border
            ws.cell(row=row, column=2, value=spec.get("qty", "N/A")).border = border
            ws.cell(row=row, column=3, value=spec.get("details", "N/A")).border = border
            ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical='top')
            row += 1
    
    # Настройка ширины колонок
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 50
    
    # Сохранение в буфер
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return buffer

